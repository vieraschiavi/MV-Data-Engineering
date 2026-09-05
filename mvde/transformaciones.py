# © 2026 Martín Viera. Todos los derechos reservados.
"""Bitácora de transformaciones: cada paso que el pipeline dio, en orden, con
tres registros (técnico, criollo, impacto río abajo) y la evidencia numérica
de la corrida. Se arma desde ``pipeline.resultados`` y el YAML, no desde una
plantilla, así lo que dice es lo que pasó. Exporta a HTML autocontenido,
Word (.docx, python-docx) y PDF (reportlab); las dos últimas bibliotecas son
opcionales y, si faltan, el exportador lo dice en vez de romper."""
from __future__ import annotations

import html as _html
import io
import json
import re
from datetime import datetime

import pandas as pd

from . import BRAND, ETAPAS, __version__
from .i18n import t

_SCD2_COLS = ("valid_from", "valid_to", "is_current", "version", "attr_hash")


# ------------------------------------------------------------------ utilidades
def _num(v, lang: str = "es") -> str:
    """Miles con punto (es/pt) o coma (en); decimales al revés."""
    if v is None or v == "—":
        return "—"
    try:
        f = float(v)
    except (TypeError, ValueError):
        return str(v)
    s = f"{f:,.0f}" if f.is_integer() else f"{f:,.2f}"
    return s if lang == "en" else s.replace(",", "X").replace(".", ",").replace("X", ".")


def _lista(items, lang: str = "es", maximo: int = 6) -> str:
    items = [str(i) for i in items]
    if not items:
        return "—"
    return ", ".join(items[:maximo]) + (f" (+{len(items) - maximo})" if len(items) > maximo else "")


def _int_en(texto: str, patron: str) -> int | None:
    m = re.search(patron, texto or "")
    return int(m.group(1).replace(".", "").replace(",", "")) if m else None


def _paso(_etapa: str, _clase: str, _objeto: str, _lang: str, _evidencia: str = "", estado: str = "ok", **d) -> dict:
    """Los posicionales llevan guion bajo para no chocar con los placeholders de las plantillas."""
    return {
        "etapa": _etapa, "tipo": _clase, "objeto": _objeto, "estado": estado, "evidencia": _evidencia,
        "tecnico": t(f"x_{_clase}_tec", _lang).format(**d),
        "criollo": t(f"x_{_clase}_cri", _lang).format(**d),
        "impacto": t(f"x_{_clase}_imp", _lang).format(**d),
    }


# ------------------------------------------------------------------ por etapa
def _fuentes(p, lang) -> list[dict]:
    r = p.resultados.get("fuentes")
    if not r or not r.ok:
        return []
    out = []
    for nombre, ev in r.evidencia.items():
        filas, cols = ev.get("filas", 0), ev.get("columnas", "—")
        out.append(_paso("fuentes", "fuente", nombre, lang, f"{_num(filas, lang)} × {cols} · {ev.get('origen', '')}",
                         nombre=nombre, tipo=ev.get("tipo", "—"), origen=ev.get("origen", "—"), filas=_num(filas, lang), columnas=cols))
    return out


def _bronze(p, lang) -> list[dict]:
    r = p.resultados.get("bronze")
    if not r or not r.ok:
        return []
    return [_paso("bronze", "bronze", n, lang, f"{_num(ev['filas'], lang)} × {ev['columnas']}", nombre=n,
                  filas=_num(ev["filas"], lang), columnas=ev["columnas"]) for n, ev in r.evidencia.items()]


def _silver(p, lang) -> list[dict]:
    r = p.resultados.get("silver")
    if not r or not r.ok:
        return []
    out = []
    cfg_all = p.spec.get("silver") or {}
    for tabla, ev in r.evidencia.items():
        cfg = cfg_all.get(tabla) or {}
        notas = ev.get("notas", [])
        n0 = len(out)
        if cfg.get("renombrar"):
            mapa = cfg["renombrar"]
            out.append(_paso("silver", "renombrar", tabla, lang, f"{len(mapa)}", tabla=tabla, n=len(mapa),
                             mapa=_lista(f"{a} → {b}" for a, b in mapa.items())))
        nota_tipos = next((n for n in notas if n.startswith("tipos:")), None)
        if nota_tipos:
            cambios = nota_tipos[len("tipos:"):].strip()
            n_c = len(cfg["tipos"]) if isinstance(cfg.get("tipos"), dict) else cambios.count(",") + 1
            out.append(_paso("silver", "tipar", tabla, lang, cambios[:120], tabla=tabla, n=n_c, cambios=cambios,
                             modo="explícito" if isinstance(cfg.get("tipos"), dict) else "auto"))
        for col, mapa in (cfg.get("decodificar") or {}).items():
            n_v = len([k for k in mapa if k != "_otros"])
            out.append(_paso("silver", "decodificar", f"{tabla}.{col}", lang, _lista(f"{k}={v}" for k, v in mapa.items()),
                             tabla=tabla, columna=col, n=n_v, otros=f", _otros → {mapa['_otros']}" if "_otros" in mapa else ""))
        if cfg.get("fechas"):
            out.append(_paso("silver", "fechas", tabla, lang, _lista(cfg["fechas"]), tabla=tabla, columnas=_lista(cfg["fechas"])))
        if cfg.get("filtrar"):
            fuera = _int_en(next((n for n in notas if n.startswith("filtro")), ""), r": (\d+) filas")
            out.append(_paso("silver", "filtrar", tabla, lang, f"{_num(fuera, lang)}", tabla=tabla, expr=cfg["filtrar"], fuera=_num(fuera, lang)))
        if cfg.get("deduplicar"):
            fuera = _int_en(next((n for n in notas if n.startswith("deduplicar")), ""), r": (\d+) duplicados")
            out.append(_paso("silver", "deduplicar", tabla, lang, f"{_num(fuera, lang)}", tabla=tabla, claves=_lista(cfg["deduplicar"]), fuera=_num(fuera, lang)))
        if cfg.get("despivotear"):
            d = cfg["despivotear"]
            filas = _int_en(next((n for n in notas if n.startswith("despivoteo")), ""), r"(\d+) filas")
            if "grupos" in d:
                detalle = f"{len(d['grupos'])} " + ("grupos de columnas" if lang == "es" else "column groups" if lang == "en" else "grupos de colunas")
            else:
                detalle = f"{len(d.get('columnas') or [])} " + ("columnas" if lang != "en" else "columns")
            out.append(_paso("silver", "despivotear", tabla, lang, f"{_num(filas, lang)}", tabla=tabla, detalle=detalle,
                             filas=_num(filas, lang), periodo=d.get("periodo") or d.get("nombre_variable") or "periodo"))
        for col, expr in (cfg.get("derivar") or {}).items():
            out.append(_paso("silver", "derivar", f"{tabla}.{col}", lang, str(expr)[:100], tabla=tabla, columna=col, expr=str(expr)))
        if cfg.get("quitar"):
            out.append(_paso("silver", "quitar", tabla, lang, _lista(cfg["quitar"]), tabla=tabla, columnas=_lista(cfg["quitar"])))
        if len(out) == n0:
            out.append(_paso("silver", "silver_nada", tabla, lang, f"{_num(ev.get('filas'), lang)} × {ev.get('columnas')}", tabla=tabla))
    return out


_RC = {"unico": "unico", "no_nulo": "no_nulo", "rango": "rango", "valores": "valores", "patron": "patron", "positivo": "positivo",
       "no_negativo": "no_negativo", "filas_min": "filas_min", "filas_exactas": "filas_exactas", "referencia": "referencia",
       "frescura": "frescura", "expresion": "expresion"}


def _calidad(p, lang) -> list[dict]:
    r = p.resultados.get("calidad")
    if not r or not p.calidad:
        return []
    out = []
    cal = p.calidad
    declaradas = list((p.spec.get("calidad") or {}).get("reglas") or [])
    for res in cal.get("resultados", []):
        tipo_r = str(res.get("regla", "")).split(":", 1)[0]     # el nombre de la regla es «tipo:tabla.columna»
        obj = f"{res['tabla']}.{res['columna']}" if res.get("columna") else res["tabla"]
        cfg = next((d for d in declaradas if d["tabla"] == res["tabla"] and d.get("columna") == res.get("columna") and d["tipo"] == tipo_r), {})
        params = {k: v for k, v in cfg.items() if k not in ("tabla", "columna", "tipo", "critico", "dimension")}
        ok = bool(res.get("paso"))
        estado = t("x_passed" if ok else "x_failed", lang)
        sev = t("x_critical" if res.get("critico") else "x_informative", lang)
        que = t(f"x_rc_{_RC.get(tipo_r, 'otro')}", lang)
        try:
            que = que.format(tipo_regla=tipo_r, a=cfg.get("a", "—"), min=cfg.get("min", "—"), max=cfg.get("max", "—"),
                             valor=_num(cfg.get("valor"), lang), dias=cfg.get("dias", 1), expresion=cfg.get("expresion", "—"))
        except (KeyError, IndexError):
            pass
        cons = t("x_regla_cons_ok" if ok else "x_regla_cons_crit" if res.get("critico") else "x_regla_cons_info", lang)
        out.append(_paso("calidad", "regla", obj, lang, res.get("detalle", ""), estado="ok" if ok else "fallo",
                         dimension=res.get("dimension", "—"), tipo_regla=tipo_r, objeto=obj, severidad=sev, estado_txt=estado,
                         detalle=res.get("detalle", ""), que_controla=que, consecuencia=cons,
                         gate=t("x_regla_gate_si" if res.get("critico") else "x_regla_gate_no", lang),
                         params=(" (" + ", ".join(f"{k}={v}" for k, v in params.items()) + ")") if params else ""))
    paso_gate = bool(cal.get("paso"))
    out.append(_paso("calidad", "gate", "calidad", lang, f"{cal.get('puntaje')}/100", estado="ok" if paso_gate else "fallo",
                     reglas=cal.get("reglas", 0), declaradas=len(declaradas), automaticas=max(0, int(cal.get("reglas", 0)) - len(declaradas)),
                     puntaje=cal.get("puntaje"), fallidas=len(cal.get("fallidas", [])), criticas=len(cal.get("criticas_fallidas", [])),
                     decision=t("x_gate_sigue" if paso_gate else "x_gate_corta", lang),
                     decision_cri=t("x_gate_sigue_cri" if paso_gate else "x_gate_corta_cri", lang)))
    return out


def _gold(p, lang) -> list[dict]:
    r = p.resultados.get("gold")
    if not r or not r.ok:
        return []
    out = []
    modelo = p.spec.get("modelo") or {}
    tablas = r.evidencia.get("tablas", {})
    dim_cfg = {d["nombre"]: d for d in modelo.get("dimensiones", []) or []}
    for d in dim_cfg.values():
        scd = int(d.get("scd", 1))
        df = p.gold.get(d["nombre"])
        attrs = d.get("atributos") or ([c for c in df.columns if c != d["clave"] and not c.endswith("_key") and c not in _SCD2_COLS] if df is not None else [])
        out.append(_paso("gold", "dimension", d["nombre"], lang, f"{_num(tablas.get(d['nombre']), lang)} · SCD {scd}",
                         nombre=d["nombre"], desde=d["desde"], clave=d["clave"], n_attr=len(attrs), attrs=_lista(attrs), scd=scd,
                         scd_detalle=t("x_scd2_detalle", lang) if scd == 2 else "", scd_cri=t("x_scd2_cri" if scd == 2 else "x_scd1_cri", lang),
                         filas=_num(tablas.get(d["nombre"]), lang)))
    if "dim_calendario" in tablas:
        nota = next((n for n in r.evidencia.get("notas", []) if n.startswith("dim_calendario")), "")
        m = re.search(r": (\S+) → (\S+)", nota)
        desde, hasta = (m.group(1), m.group(2)) if m else ("—", "—")
        cal = modelo.get("calendario", "auto")
        out.append(_paso("gold", "calendario", "dim_calendario", lang, f"{_num(tablas['dim_calendario'], lang)} días",
                         modo=t("x_calendario_auto" if cal == "auto" else "x_calendario_decl", lang), desde=desde, hasta=hasta,
                         filas=_num(tablas["dim_calendario"], lang)))
    for h in modelo.get("hechos", []) or []:
        claves = h.get("claves") or {}
        joins = "; ".join(t("x_join_txt", lang).format(col=c, dim=dm) for c, dm in claves.items()) or "—"
        out.append(_paso("gold", "hecho", h["nombre"], lang, f"{_num(tablas.get(h['nombre']), lang)}",
                         nombre=h["nombre"], desde=h["desde"], joins=joins,
                         fecha=t("x_fecha_key_txt", lang).format(col=h["fecha"]) if h.get("fecha") else "",
                         medidas=_lista(h.get("medidas") or []), filas=_num(tablas.get(h["nombre"]), lang),
                         dims=_lista(claves.values()) + (" + dim_calendario" if h.get("fecha") else "")))
    for n in tablas:
        if n.startswith("tbl_"):
            df = p.gold.get(n)
            out.append(_paso("gold", "tabla_directa", n, lang, f"{_num(tablas[n], lang)}", nombre=n, filas=_num(tablas[n], lang),
                             fk=t("x_fk_auto", lang) if df is not None and "fecha_key" in df.columns else ""))
    return out


def _almacen(p, lang) -> list[dict]:
    r = p.resultados.get("almacen")
    if not r or not r.ok:
        return []
    ev = r.evidencia
    out = [_paso("almacen", "almacen", "duckdb", lang, _lista(ev.get("tablas", {}), lang, 8), n=len(ev.get("tablas", {})), ruta=p.ruta_db.name)]
    vistas = p.spec.get("vistas") or {}
    for v in ev.get("vistas", []):
        sql = " ".join(str(vistas.get(v, "")).split())
        out.append(_paso("almacen", "vista", v, lang, sql[:90] + ("…" if len(sql) > 90 else ""), nombre=v, sql=sql))
    pub = (p.spec.get("almacen") or {}).get("publicar") or {}
    if ev.get("publicacion"):
        url = re.sub(r"://[^@]+@", "://***@", str(pub.get("url", "")))
        out.append(_paso("almacen", "publicar", "publicacion", lang, url, url=url, esquema=pub.get("esquema") or "gold"))
    return out


def _gobernanza(p, lang) -> list[dict]:
    r = p.resultados.get("gobernanza")
    if not r or not r.ok:
        return []
    pj = r.evidencia.get("puntajes", {})
    cat = p.catalogo if isinstance(p.catalogo, pd.DataFrame) else pd.DataFrame()
    out = [_paso("gobernanza", "catalogo", "catalogo", lang, f"{pj.get('tablas', '—')} / {len(cat)} · doc {pj.get('documentacion_pct', '—')} %",
                 tablas=pj.get("tablas", "—"), columnas=len(cat), documentacion=pj.get("documentacion_pct", "—"))]
    n_ar = len(r.evidencia.get("linaje", []))
    out.append(_paso("gobernanza", "linaje", "linaje", lang, f"{n_ar}", aristas=n_ar))
    pii_cols = list(cat[cat["pii"]].apply(lambda x: f"{x['tabla']}.{x['columna']}", axis=1).unique()) if len(cat) and "pii" in cat.columns else []
    decl = (p.spec.get("gobernanza") or {}).get("pii") or []
    if pii_cols:
        out.append(_paso("gobernanza", "pii", "pii", lang, _lista(pii_cols), n=len(pii_cols), columnas=_lista(pii_cols), declaradas=len(decl)))
    else:
        s = _paso("gobernanza", "pii", "pii", lang, "0", n=0, columnas="—", declaradas=len(decl))
        s["criollo"] = t("x_pii_ninguna", lang)
        out.append(s)
    return out


def _ml(p, lang) -> list[dict]:
    r = p.resultados.get("ml")
    if not r or not r.ok or r.omitida or not p.ml:
        return []
    cfg, m = p.spec.get("ml") or {}, p.ml
    notas = m.get("notas", [])
    met = m.get("metricas") or {}
    fuga = _int_en(next((n for n in notas if n.startswith("fuga")), ""), r"quitadas (\d+)")
    n_feat = len(m.get("importancia") or {})
    total = int(m.get("filas_train", 0)) + int(m.get("filas_seleccion", 0)) + int(m.get("filas_holdout", 0))
    out = [_paso("ml", "ml_prep", cfg.get("target", "—"), lang, f"{_num(total, lang)} · {cfg.get('target')}",
                 origen=t("x_ml_origen_sql", lang) if cfg.get("sql") else t("x_ml_origen_tabla", lang).format(tabla=cfg.get("tabla", "—")),
                 filas=_num(total, lang), target=cfg.get("target", "—"), tipo=m.get("tipo", "—"), n_feat=f"≥{n_feat}" if n_feat else "—",
                 excluidas=(" y excluidas " + _lista(cfg["excluir"])) if cfg.get("excluir") else "",
                 fuga=t("x_ml_fuga_si", lang).format(n=fuga) if fuga else t("x_ml_fuga_no", lang))]
    out.append(_paso("ml", "ml_corte", "60/20/20", lang, f"{m.get('filas_train')} / {m.get('filas_seleccion')} / {m.get('filas_holdout')}",
                     train=_num(m.get("filas_train"), lang), sel=_num(m.get("filas_seleccion"), lang), ho=_num(m.get("filas_holdout"), lang),
                     modo=t("x_ml_corte_temporal", lang).format(fecha=cfg["fecha"]) if cfg.get("fecha") else t("x_ml_corte_aleatorio", lang)))
    comp = m.get("comparacion") or []
    clave = "sel_auc" if m.get("tipo") == "clasificacion" else "sel_r2"
    tabla = "; ".join(f"{c['modelo']} {clave[4:].upper()} {c.get(clave)}" for c in comp) or "—"
    out.append(_paso("ml", "ml_modelos", m.get("modelo", "—"), lang, tabla, n=len(comp), tabla=tabla))
    brecha = m.get("brecha_seleccion_holdout")
    brecha_txt = _num(brecha, lang) if brecha is not None else "—"
    metricas = " · ".join(f"{k} {v}" for k, v in met.items() if v is not None) or "—"
    if m.get("tipo") == "clasificacion":
        metrica_cri = t("x_ml_auc_cri", lang).format(auc_pct=round(100 * float(met.get("auc") or 0)))
    else:
        metrica_cri = t("x_ml_r2_cri", lang).format(r2_pct=round(100 * float(met.get("r2") or 0)), mae=_num(met.get("mae"), lang))
    mal = brecha is not None and abs(float(brecha)) > 0.05
    s = _paso("ml", "ml_elegido", m.get("modelo", "—"), lang, metricas, modelo=m.get("modelo", "—"), metricas=metricas, brecha=brecha_txt,
              alerta=t("x_ml_brecha_alerta", lang) if mal else "", metrica_cri=metrica_cri,
              brecha_cri=t("x_ml_brecha_mal" if mal else "x_ml_brecha_ok", lang).format(brecha=brecha_txt),
              lift=_num(met.get("lift_decil10"), lang))
    if m.get("tipo") != "clasificacion":
        s["impacto"] = t("x_ml_elegido_imp_reg", lang).format(mae=_num(met.get("mae"), lang))
    out.append(s)
    imp = m.get("importancia") or {}
    if imp:
        top = list(imp.items())[:6]
        out.append(_paso("ml", "ml_importancia", "importancia", lang, _lista(f"{k} {v}" for k, v in top),
                         top=_lista(f"{k} ({v})" for k, v in top), top_cri=_lista((k for k, _ in top), lang, 4)))
    kobra = "estrategia" in (p.gold.get("ml_scores", pd.DataFrame()).columns)
    out.append(_paso("ml", "ml_score", "ml_scores", lang, f"{_num(m.get('filas_scoreadas'), lang)}", filas=_num(m.get("filas_scoreadas"), lang),
                     objetivo=t("x_ml_score_otro" if (cfg.get("tabla_score") or cfg.get("sql_score")) else "x_ml_score_self", lang),
                     kobra=t("x_ml_kobra_tec", lang) if kobra else "", kobra_cri=t("x_ml_kobra_cri", lang) if kobra else ""))
    return out


def _reporte(p, lang) -> list[dict]:
    r = p.resultados.get("reporte")
    if not r or not r.ok:
        return []
    out = []
    auto = t("x_kpi_auto", lang) if p.spec.get("_kpis_automaticos") else ""
    specs = {k["nombre"]: k for k in p.spec.get("kpis") or []}
    for k in p.kpis or []:
        cfg = specs.get(k["nombre"], {})
        tipo = cfg.get("tipo", "agregacion")
        if tipo == "ratio":
            n, d = cfg.get("numerador", {}), cfg.get("denominador", {})
            fmt = lambda x: f"{t('x_agg_' + x.get('agregacion', 'sum'), lang)} " + (f"«{x.get('columna')}»" if x.get("columna") else f"«{x.get('tabla')}»")  # noqa: E731
            como = t("x_kpi_como_ratio", lang).format(num=fmt(n), den=fmt(d))
        elif tipo == "sql":
            como = t("x_kpi_como_sql", lang)
        elif cfg.get("agregacion") == "count":
            como = t("x_kpi_como_count", lang).format(tabla=cfg.get("tabla", "—"))
        else:
            como = t("x_kpi_como_agg", lang).format(agg_cri=t("x_agg_" + cfg.get("agregacion", "sum"), lang), columna=cfg.get("columna", "—"), tabla=cfg.get("tabla", "—"))
        valor = k["texto"] if k.get("ok") else t("x_kpi_error", lang).format(error=k.get("error", ""))
        out.append(_paso("reporte", "kpi", k["nombre"], lang, valor, estado="ok" if k.get("ok") else "fallo",
                         nombre=k["nombre"], sql=" ".join(str(k.get("sql", "")).split()), valor=valor,
                         formato=f" · {cfg['formato']}" if cfg.get("formato") else "", por=f" · por {cfg['por']}" if cfg.get("por") else "",
                         como_cri=como + auto))
    for g in r.evidencia.get("graficos", []):
        nombre = g.rsplit(".", 1)[0]
        regla = t("x_grafico_regla_serie" if "serie" in nombre or "mensual" in nombre else "x_grafico_regla_por", lang)
        out.append(_paso("reporte", "grafico", nombre, lang, g, nombre=nombre, regla=regla))
    out.append(_paso("reporte", "reporte_arch", "reporte.xlsx / reporte.html", lang, "xlsx + html"))
    return out


def _dax(p, lang) -> list[dict]:
    r = p.resultados.get("dax")
    if not r or not r.ok:
        return []
    out = []
    for m in p.medidas or []:
        n = m["nombre"]
        clave = ("x_medida_var" if "var." in n else "x_medida_prev" if "anterior" in n or "previous" in n else "x_medida_pct" if "% del total" in n or "% of total" in n
                 else "x_medida_pp" if "vs total" in n else "x_medida_total" if n.endswith(" total") else "x_medida_filtrado" if "(filtrado)" in n else "x_medida_base")
        out.append(_paso("dax", "medida", n, lang, m.get("formato", ""), nombre=n, tabla=m.get("tabla", "—"),
                         expresion=" ".join(str(m.get("expresion", "")).split()), formato=m.get("formato") or "—", que_cri=t(clave, lang)))
    return out


def _powerbi(p, lang) -> list[dict]:
    r = p.resultados.get("powerbi")
    if not r or not r.ok or r.omitida:
        return []
    ev = r.evidencia
    hall = ev.get("hallazgos") or []
    s = _paso("powerbi", "pbi", "pbit / pbip", lang, f"auditoría {ev.get('auditoria')}/100", tablas=ev.get("tablas", "—"), relaciones=ev.get("relaciones", "—"),
              medidas=ev.get("medidas", "—"), auditoria=ev.get("auditoria", "—"),
              hallazgos=("; " + _lista(f"{a} {b}" for a, b in hall)) if hall else "")
    if ev.get("demo_embebida") is not None:
        s["criollo"] += t("x_pbi_demo", lang)
    return [s]


def _entrega(p, lang) -> list[dict]:
    r = p.resultados.get("entrega")
    if not r or not r.ok:
        return []
    salud = (r.evidencia.get("salud") or {}).get("total", "—")
    carpeta = p.dirs["entrega"]
    n_arch = len(list(carpeta.iterdir())) if carpeta.exists() else 0
    idiomas = " + ".join(dict.fromkeys([p.spec.get("idioma", "es"), "en"]))
    return [_paso("entrega", "entrega", "entrega/", lang, f"{n_arch} · salud {salud}", idiomas=idiomas, n_arch=n_arch, salud=salud)]


_CONSTRUCTORES = {"fuentes": _fuentes, "bronze": _bronze, "silver": _silver, "calidad": _calidad, "gold": _gold, "almacen": _almacen,
                  "gobernanza": _gobernanza, "ml": _ml, "reporte": _reporte, "dax": _dax, "powerbi": _powerbi, "entrega": _entrega}


# ------------------------------------------------------------------ API
def pasos(pipeline, lang: str = "es") -> list[dict]:
    """La secuencia completa, numerada, en el orden de las etapas."""
    out = []
    for e in ETAPAS:
        for s in _CONSTRUCTORES[e](pipeline, lang):
            s["n"] = len(out) + 1
            s["etapa_titulo"] = t(f"st_{e}", lang)
            out.append(s)
    return out


def resumen(pasos_: list[dict], lang: str = "es") -> list[dict]:
    filas = []
    for e in ETAPAS:
        de = [s for s in pasos_ if s["etapa"] == e]
        if de:
            filas.append({"etapa": t(f"st_{e}", lang), "pasos": len(de), "fallos": sum(s["estado"] == "fallo" for s in de),
                          "objetos": _lista(dict.fromkeys(s["objeto"] for s in de), lang, 5)})
    return filas


def _cabecera(pipeline, pasos_, lang) -> dict:
    fecha = datetime.now().strftime("%Y-%m-%d %H:%M")
    return {"proyecto": pipeline.spec["nombre"], "titulo": t("x_titulo", lang), "fecha": fecha,
            "intro": t("x_intro", lang).format(fecha=fecha, n_pasos=len(pasos_), n_etapas=len({s['etapa'] for s in pasos_})),
            "pie": t("x_generated", lang).format(version=__version__, fecha=fecha)}


def html(pipeline, lang: str = "es") -> str:
    """HTML autocontenido con el estilo de marca y un selector técnico / criollo / todo."""
    ps = pasos(pipeline, lang)
    cab = _cabecera(pipeline, ps, lang)
    e = _html.escape
    filas_res = "".join(f"<tr><td>{e(r['etapa'])}</td><td>{r['pasos']}</td><td>{r['fallos']}</td><td>{e(r['objetos'])}</td></tr>" for r in resumen(ps, lang))
    cuerpo, etapa_actual = [], None
    for s in ps:
        if s["etapa"] != etapa_actual:
            etapa_actual = s["etapa"]
            cuerpo.append(f"<h2 id='{etapa_actual}'>{e(s['etapa_titulo'])}</h2><p class='d'>{e(t('d_' + etapa_actual, lang))}</p>")
        cuerpo.append(
            f"<article class='paso {s['estado']}'><div class='n'>{s['n']}</div><div class='b'>"
            f"<div class='o'>{e(s['objeto'])} <small>{e(s['tipo'])}</small></div>"
            f"<p class='tec'><b>{e(t('x_tec', lang))}</b> {e(s['tecnico'])}</p>"
            f"<p class='cri'><b>{e(t('x_cri', lang))}</b> {e(s['criollo'])}</p>"
            f"<p class='imp'><b>{e(t('x_imp', lang))}</b> {e(s['impacto'])}</p>"
            f"<p class='ev'>{e(t('x_ev', lang))}: {e(str(s['evidencia']))}</p></div></article>")
    return f"""<!doctype html><html lang="{lang}"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>{e(cab['proyecto'])} · {e(cab['titulo'])}</title>
<style>
body{{font-family:Segoe UI,Inter,Arial,sans-serif;margin:0;background:#f8fafc;color:#111827}}
header{{background:linear-gradient(160deg,{BRAND['navy']},#0a1a2f);color:{BRAND['ink']};padding:28px 40px}}
header h1{{margin:0;font-size:24px}} header p{{margin:8px 0 0;color:{BRAND['muted']};max-width:900px}}
.badge{{display:inline-block;background:rgba(242,180,65,.15);border:1px solid rgba(242,180,65,.5);color:{BRAND['amber']};border-radius:20px;padding:3px 12px;font-size:12px;font-weight:700;letter-spacing:.06em;text-transform:uppercase}}
main{{padding:24px 40px;max-width:1100px}} h2{{color:{BRAND['navy']};margin:28px 0 4px;border-bottom:2px solid {BRAND['amber']};padding-bottom:4px}}
p.d{{color:#6b7280;margin:0 0 12px;font-size:14px}}
table{{border-collapse:collapse;width:100%;background:#fff;margin:12px 0}} th{{background:{BRAND['navy']};color:#fff;text-align:left;padding:8px}} td{{padding:8px;border-bottom:1px solid #e5e7eb;vertical-align:top}}
.paso{{display:flex;gap:14px;background:#fff;border:1px solid #e5e7eb;border-left:4px solid {BRAND['green']};border-radius:10px;padding:12px 16px;margin:8px 0;box-shadow:0 1px 3px rgba(0,0,0,.05)}}
.paso.fallo{{border-left-color:{BRAND['red']}}} .n{{font-weight:700;color:{BRAND['amber']};font-size:20px;min-width:34px}} .b{{flex:1}}
.o{{font-weight:700;color:{BRAND['navy']};margin-bottom:4px}} .o small{{color:#9ca3af;font-weight:400;margin-left:6px}}
.b p{{margin:4px 0;font-size:14px;line-height:1.45}} .b p b{{color:{BRAND['navy']};margin-right:4px}} p.ev{{color:#6b7280;font-size:12.5px}}
.cri b{{color:{BRAND['amber2']}!important}} .imp b{{color:{BRAND['blue']}!important}}
.ctl{{position:sticky;top:0;background:#f8fafc;padding:10px 0;border-bottom:1px solid #e5e7eb;z-index:2}} .ctl button{{border:1px solid {BRAND['navy']};background:#fff;color:{BRAND['navy']};border-radius:20px;padding:6px 14px;margin-right:6px;cursor:pointer;font-weight:600}}
.ctl button.on{{background:{BRAND['navy']};color:#fff}} body.solo-tec p.cri,body.solo-tec p.imp{{display:none}} body.solo-cri p.tec,body.solo-cri p.ev{{display:none}}
footer{{color:#6b7280;font-size:12px;padding:20px 40px}}
@media print{{.ctl{{display:none}} .paso{{break-inside:avoid}} header{{-webkit-print-color-adjust:exact;print-color-adjust:exact}}}}
</style></head><body>
<header><span class="badge">MV · Data Engineering</span><h1>{e(cab['proyecto'])} · {e(cab['titulo'])}</h1><p>{e(cab['intro'])}</p></header>
<main>
<div class="ctl"><b>{e(t('x_view', lang))}:</b>
<button class="on" onclick="v('',this)">{e(t('x_view_all', lang))}</button><button onclick="v('solo-tec',this)">{e(t('x_view_tec', lang))}</button><button onclick="v('solo-cri',this)">{e(t('x_view_cri', lang))}</button></div>
<h2>{e(t('x_summary', lang))}</h2>
<table><tr><th>{e(t('x_etapa', lang))}</th><th>{e(t('x_n_pasos', lang))}</th><th>{e(t('x_failed', lang))}</th><th>{e(t('x_objeto', lang))}</th></tr>{filas_res}</table>
{''.join(cuerpo)}
</main><footer>{e(cab['pie'])}</footer>
<script>function v(c,b){{document.body.className=c;document.querySelectorAll('.ctl button').forEach(x=>x.classList.remove('on'));b.classList.add('on');}}</script>
</body></html>"""


def docx(pipeline, lang: str = "es") -> bytes:
    """Word con la misma estructura. Requiere python-docx."""
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: F401
        from docx.shared import Pt, RGBColor
    except ImportError as exc:  # pragma: no cover - depende de la instalación
        raise RuntimeError(t("x_lib_missing", lang).format(lib="python-docx")) from exc
    ps = pasos(pipeline, lang)
    cab = _cabecera(pipeline, ps, lang)
    navy, amber, blue, gris = RGBColor(0x08, 0x15, 0x27), RGBColor(0xe3, 0x9a, 0x2e), RGBColor(0x2f, 0x74, 0xc0), RGBColor(0x6b, 0x72, 0x80)
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name, st.font.size = "Segoe UI", Pt(10)
    doc.add_heading(f"{cab['proyecto']} · {cab['titulo']}", 0)
    doc.add_paragraph(cab["intro"])
    doc.add_heading(t("x_summary", lang), 1)
    tabla = doc.add_table(rows=1, cols=4)
    tabla.style = "Light Grid Accent 1"
    for i, h in enumerate((t("x_etapa", lang), t("x_n_pasos", lang), t("x_failed", lang), t("x_objeto", lang))):
        tabla.rows[0].cells[i].text = h
    for r in resumen(ps, lang):
        c = tabla.add_row().cells
        c[0].text, c[1].text, c[2].text, c[3].text = r["etapa"], str(r["pasos"]), str(r["fallos"]), r["objetos"]
    etapa_actual = None
    for s in ps:
        if s["etapa"] != etapa_actual:
            etapa_actual = s["etapa"]
            doc.add_heading(s["etapa_titulo"], 1)
            doc.add_paragraph(t("d_" + etapa_actual, lang)).runs[0].font.color.rgb = gris
        h = doc.add_heading(f"{t('x_paso', lang)} {s['n']} · {s['objeto']}", 2)
        if s["estado"] == "fallo":
            h.runs[0].font.color.rgb = RGBColor(0xe0, 0x5c, 0x5c)
        for clave, color, texto in (("x_tec", navy, s["tecnico"]), ("x_cri", amber, s["criollo"]), ("x_imp", blue, s["impacto"])):
            par = doc.add_paragraph()
            run = par.add_run(t(clave, lang) + ": ")
            run.bold, run.font.color.rgb = True, color
            par.add_run(texto)
        ev = doc.add_paragraph().add_run(f"{t('x_ev', lang)}: {s['evidencia']}")
        ev.italic, ev.font.size, ev.font.color.rgb = True, Pt(8.5), gris
    pie = doc.add_paragraph().add_run(cab["pie"])
    pie.font.size, pie.font.color.rgb = Pt(8), gris
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def pdf(pipeline, lang: str = "es") -> bytes:
    """PDF con reportlab (platypus). Usa DejaVu Sans si está (la trae matplotlib)
    para que flechas, comillas y acentos se vean; si no, Helvetica."""
    try:
        from reportlab.lib import colors
        from reportlab.lib.enums import TA_LEFT
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import mm
        from reportlab.platypus import KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    except ImportError as exc:  # pragma: no cover - depende de la instalación
        raise RuntimeError(t("x_lib_missing", lang).format(lib="reportlab")) from exc
    fuente, fuente_b = _fuente_pdf()
    ps = pasos(pipeline, lang)
    cab = _cabecera(pipeline, ps, lang)
    base = getSampleStyleSheet()
    S = {
        "t": ParagraphStyle("t", parent=base["Title"], fontName=fuente_b, fontSize=17, textColor=colors.HexColor(BRAND["navy"]), alignment=TA_LEFT, spaceAfter=4),
        "h1": ParagraphStyle("h1", parent=base["Heading1"], fontName=fuente_b, fontSize=13, textColor=colors.HexColor(BRAND["navy"]), spaceBefore=12, spaceAfter=2),
        "d": ParagraphStyle("d", parent=base["Normal"], fontName=fuente, fontSize=8.5, textColor=colors.HexColor("#6b7280"), spaceAfter=6),
        "o": ParagraphStyle("o", parent=base["Normal"], fontName=fuente_b, fontSize=10, textColor=colors.HexColor(BRAND["navy"]), spaceBefore=6),
        "p": ParagraphStyle("p", parent=base["Normal"], fontName=fuente, fontSize=9, leading=12),
        "ev": ParagraphStyle("ev", parent=base["Normal"], fontName=fuente, fontSize=7.5, textColor=colors.HexColor("#6b7280"), spaceAfter=4),
    }
    esc = lambda x: _html.escape(str(x))  # noqa: E731
    story = [Paragraph(f"{esc(cab['proyecto'])} · {esc(cab['titulo'])}", S["t"]), Paragraph(esc(cab["intro"]), S["p"]), Spacer(1, 6),
             Paragraph(esc(t("x_summary", lang)), S["h1"])]
    datos = [[t("x_etapa", lang), t("x_n_pasos", lang), t("x_failed", lang), t("x_objeto", lang)]] + \
            [[r["etapa"], str(r["pasos"]), str(r["fallos"]), Paragraph(esc(r["objetos"]), S["ev"])] for r in resumen(ps, lang)]
    tb = Table(datos, colWidths=[45 * mm, 18 * mm, 18 * mm, 95 * mm])
    tb.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(BRAND["navy"])), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                            ("FONTNAME", (0, 0), (-1, -1), fuente), ("FONTSIZE", (0, 0), (-1, -1), 8), ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#e5e7eb")),
                            ("VALIGN", (0, 0), (-1, -1), "TOP")]))
    story.append(tb)
    etapa_actual = None
    col = {"x_tec": BRAND["navy"], "x_cri": BRAND["amber2"], "x_imp": BRAND["blue"]}
    for s in ps:
        if s["etapa"] != etapa_actual:
            etapa_actual = s["etapa"]
            story += [Paragraph(esc(s["etapa_titulo"]), S["h1"]), Paragraph(esc(t("d_" + etapa_actual, lang)), S["d"])]
        color_o = BRAND["red"] if s["estado"] == "fallo" else BRAND["navy"]
        bloque = [Paragraph(f"<font color='{BRAND['amber']}'>{s['n']}</font>&nbsp;&nbsp;<font color='{color_o}'>{esc(s['objeto'])}</font> <font color='#9ca3af' size='7'>{esc(s['tipo'])}</font>", S["o"])]
        for clave, texto in (("x_tec", s["tecnico"]), ("x_cri", s["criollo"]), ("x_imp", s["impacto"])):
            bloque.append(Paragraph(f"<font color='{col[clave]}'><b>{esc(t(clave, lang))}:</b></font> {esc(texto)}", S["p"]))
        bloque.append(Paragraph(f"{esc(t('x_ev', lang))}: {esc(s['evidencia'])}", S["ev"]))
        story.append(KeepTogether(bloque))
    story += [Spacer(1, 10), Paragraph(esc(cab["pie"]), S["ev"])]
    buf = io.BytesIO()
    SimpleDocTemplate(buf, pagesize=A4, leftMargin=16 * mm, rightMargin=16 * mm, topMargin=16 * mm, bottomMargin=16 * mm,
                      title=f"{cab['proyecto']} · {cab['titulo']}", author="MV Data Engineering").build(story)
    return buf.getvalue()


def _fuente_pdf() -> tuple[str, str]:
    try:
        from matplotlib import font_manager as fm
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        regular = fm.findfont("DejaVu Sans", fallback_to_default=False)
        negrita = fm.findfont(fm.FontProperties(family="DejaVu Sans", weight="bold"), fallback_to_default=False)
        pdfmetrics.registerFont(TTFont("DejaVu", regular))
        pdfmetrics.registerFont(TTFont("DejaVu-Bold", negrita))
        return "DejaVu", "DejaVu-Bold"
    except Exception:  # noqa: BLE001 - sin DejaVu se cae a Helvetica, que alcanza
        return "Helvetica", "Helvetica-Bold"


def exportar(pipeline, carpeta, lang: str = "es") -> dict[str, str]:
    """Escribe TRANSFORMACIONES_<lang>.{json,html,docx,pdf} en `carpeta`; devuelve
    {formato: ruta} y, si una biblioteca falta, {formato: 'error: ...'} sin romper."""
    from pathlib import Path
    carpeta = Path(carpeta)
    carpeta.mkdir(parents=True, exist_ok=True)
    ps = pasos(pipeline, lang)
    salida = {}
    p = carpeta / f"TRANSFORMACIONES_{lang}.json"
    p.write_text(json.dumps(ps, indent=2, ensure_ascii=False), encoding="utf-8")
    salida["json"] = str(p)
    p = carpeta / f"TRANSFORMACIONES_{lang}.html"
    p.write_text(html(pipeline, lang), encoding="utf-8")
    salida["html"] = str(p)
    for fmt, fn in (("docx", docx), ("pdf", pdf)):
        try:
            p = carpeta / f"TRANSFORMACIONES_{lang}.{fmt}"
            p.write_bytes(fn(pipeline, lang))
            salida[fmt] = str(p)
        except RuntimeError as exc:
            salida[fmt] = f"error: {exc}"
    return salida
