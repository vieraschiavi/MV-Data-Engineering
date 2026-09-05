"""Bitácora de transformaciones: secuencia completa en orden, tres registros por
paso (técnico / criollo / impacto), paridad de idiomas y exportación HTML / Word / PDF."""
import io
import json

import pytest

from mvde import ETAPAS, demos, transformaciones
from mvde.orquestador import Pipeline


@pytest.fixture(scope="module")
def cob(tmp_path_factory):
    ruta = demos.crear("cobranzas", tmp_path_factory.mktemp("cob"))
    p = Pipeline.desde_yaml(ruta)
    p.correr()
    return p


def test_pasos_cubren_todas_las_etapas_en_orden(cob):
    ps = transformaciones.pasos(cob, "es")
    assert [s["n"] for s in ps] == list(range(1, len(ps) + 1))
    etapas = list(dict.fromkeys(s["etapa"] for s in ps))
    assert etapas == [e for e in ETAPAS if e in etapas] and set(etapas) == set(ETAPAS)
    tipos = {s["tipo"] for s in ps}
    assert {"fuente", "bronze", "decodificar", "deduplicar", "derivar", "regla", "gate", "dimension", "calendario", "hecho",
            "almacen", "vista", "catalogo", "linaje", "pii", "ml_prep", "ml_corte", "ml_modelos", "ml_elegido", "ml_score",
            "kpi", "grafico", "medida", "pbi", "entrega"} <= tipos


def test_cada_paso_tiene_los_tres_registros_sin_placeholders(cob):
    for lang in ("es", "en", "pt"):
        ps = transformaciones.pasos(cob, lang)
        for s in ps:
            for k in ("tecnico", "criollo", "impacto"):
                assert s[k] and "{" not in s[k] and "}" not in s[k], (lang, s["n"], k, s[k])
    assert len(transformaciones.pasos(cob, "es")) == len(transformaciones.pasos(cob, "en")) == len(transformaciones.pasos(cob, "pt"))


def test_usa_los_numeros_reales_de_la_corrida(cob):
    ps = transformaciones.pasos(cob, "es")
    fuente = next(s for s in ps if s["tipo"] == "fuente" and s["objeto"] == "cuotas")
    assert "48.000" in fuente["tecnico"] and "48.000" in fuente["criollo"]
    regla_mala = next(s for s in ps if s["tipo"] == "regla" and s["estado"] == "fallo")
    assert "no cumple" in regla_mala["criollo"] and "no_negativo" in regla_mala["tecnico"]
    gate = next(s for s in ps if s["tipo"] == "gate")
    assert str(cob.calidad["puntaje"]) in gate["tecnico"] and "sigue" in gate["tecnico"]
    dim = next(s for s in ps if s["tipo"] == "dimension")
    assert "SCD tipo 2" in dim["tecnico"] and "no se pisa" in dim["criollo"]
    ml = next(s for s in ps if s["tipo"] == "ml_elegido")
    assert cob.ml["modelo"] in ml["tecnico"] and str(cob.ml["metricas"]["auc"]) in ml["tecnico"]
    en = transformaciones.pasos(cob, "en")
    assert "48,000" in next(s for s in en if s["tipo"] == "fuente" and s["objeto"] == "cuotas")["tecnico"]


def test_exporta_html_word_pdf_y_json(cob):
    ps = transformaciones.pasos(cob, "es")
    h = transformaciones.html(cob, "es")
    assert h.count("<article class='paso") == len(ps) and "Sólo criollo" in h and "<script>" in h
    docx_bytes = transformaciones.docx(cob, "es")
    from docx import Document
    doc = Document(io.BytesIO(docx_bytes))
    assert sum("Paso " in par.text for par in doc.paragraphs) == len(ps)
    pdf_bytes = transformaciones.pdf(cob, "es")
    assert pdf_bytes[:5] == b"%PDF-" and len(pdf_bytes) > 20_000
    entrega = cob.dirs["entrega"]
    for lang in ("es", "en"):
        for fmt in ("json", "html", "docx", "pdf"):
            assert (entrega / f"TRANSFORMACIONES_{lang}.{fmt}").exists(), (lang, fmt)
    assert len(json.loads((entrega / "TRANSFORMACIONES_en.json").read_text(encoding="utf-8"))) == len(ps)
    assert cob.transformaciones["es"]["pdf"].endswith(".pdf")


def test_resumen_por_etapa(cob):
    ps = transformaciones.pasos(cob, "pt")
    res = transformaciones.resumen(ps, "pt")
    assert sum(r["pasos"] for r in res) == len(ps) and sum(r["fallos"] for r in res) == sum(s["estado"] == "fallo" for s in ps)
    assert res[0]["etapa"].endswith("Fontes")
